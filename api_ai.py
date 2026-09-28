"""
api_ai.py — the grounded generation endpoint.

Separate from the deterministic pipeline on purpose. If this file is removed,
the rest of the product still works and still produces a complete structure.
That is the fallback behaviour a teacher-facing tool needs.
"""

from __future__ import annotations

import json
import re
import uuid
from typing import Any

from fastapi import APIRouter, Body, Header, HTTPException

import clil_engine
import generator
import grounding

router = APIRouter()
_db = None


def attach(db) -> APIRouter:
    global _db
    _db = db
    return router


def vid(visitor: str | None) -> str:
    return visitor or "local-dev-visitor"


@router.post("/api/generate/grounded", status_code=201)
def generate_grounded(body: dict[str, Any] = Body(...),
                      x_visitor_id: str | None = Header(default=None)) -> dict[str, Any]:
    """Build the structure deterministically, then ground it, then fill it.

    Order matters. The structure exists before any model is called, so a model
    failure degrades to a complete skeleton rather than to nothing.
    """
    visitor = vid(x_visitor_id)
    course_id = body.get("course_id")

    if course_id:
        row = _db.execute("SELECT data FROM courses WHERE id = ? AND visitor_id = ?",
                          [course_id, visitor]).fetchone()
        if row:
            course = json.loads(row["data"])
            for key in ("subject", "target_language", "first_language", "learner_group",
                        "support_level", "purpose", "setting", "format", "teacher_role",
                        "minutes", "lesson_total"):
                body.setdefault(key, course.get(key))
            body.setdefault("course_name", course.get("name"))
            body["mode"] = "course"
            taught = _db.execute("SELECT COUNT(*) c FROM lessons WHERE course_id = ?",
                                 [course_id]).fetchone()["c"]
            body.setdefault("lesson_number", taught + 1)

    package = clil_engine.build_package(body)

    # Retrieval may be supplied by the caller. Useful when search runs in a
    # different process to the web server, which is how the exemplar builder
    # works, and how a production deployment would separate the two concerns.
    evidence = body.pop("evidence", None)
    if not evidence:
        try:
            evidence = grounding.build_evidence(
                body.get("subject") or "", body.get("topic") or "",
                limit=int(body.get("source_limit") or 6),
            )
        except Exception as exc:                              # retrieval is the fragile step
            evidence = {"sources": [], "corpus": "", "searched": 0, "usable": 0,
                        "error": f"{type(exc).__name__}: {exc}"}

    try:
        package = generator.enrich(package, body, evidence)
        if evidence.get("error"):
            package.setdefault("generation", {})["retrieval_error"] = evidence["error"]
    except Exception as exc:
        package = clil_engine.enrich_with_model(package, body)
        package["generation"]["error"] = (
            f"Grounded generation failed ({type(exc).__name__}: {exc}). "
            f"The structure below is complete; the subject detail is yours to write."
        )

    token = uuid.uuid4().hex
    _db.execute(
        "INSERT INTO lessons (id, visitor_id, course_id, share_token, payload, package) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        [package["id"], visitor, course_id, token, json.dumps(body), json.dumps(package)],
    )
    _db.commit()
    package["share_token"] = token
    return package


@router.post("/api/audit")
def audit(body: dict[str, Any] = Body(...)) -> dict[str, Any]:
    """Re-run the audit pass over an existing package.

    Auditing is far cheaper than generating, so a teacher can re-check a
    package they have edited without paying to rebuild it.
    """
    package = body.get("package")
    if not package:
        raise HTTPException(400, "No package supplied")
    evidence = body.get("evidence") or {"corpus": "", "sources": []}
    return generator.reaudit(package, body.get("payload") or {}, evidence)


@router.get("/api/exemplars")
def list_exemplars() -> list[dict[str, Any]]:
    """Published exemplar packages, served as static content.

    Exemplars exist so a teacher can judge the quality of the output without
    anyone spending anything on generation.
    """
    import os
    folder = os.path.join(os.path.dirname(os.path.abspath(__file__)), "exemplars")
    if not os.path.isdir(folder):
        return []
    out = []
    for name in sorted(os.listdir(folder)):
        # Files beginning with an underscore are working files, such as the
        # cached retrieval corpus. They are not exemplars.
        if not name.endswith(".json") or name.startswith("_"):
            continue
        with open(os.path.join(folder, name), encoding="utf-8") as fh:
            pkg = json.load(fh)
        out.append({
            "slug": name[:-5],
            "title": pkg["header"]["title"],
            "subject": pkg["header"]["subject"],
            "learner_group": pkg["header"]["learner_group"],
            "teacher_role": pkg["header"]["teacher_role"],
            "sources": len((pkg.get("grounded") or {}).get("sources", [])),
        })
    return out


@router.get("/api/exemplars/{slug}")
def get_exemplar(slug: str) -> dict[str, Any]:
    import os
    safe = "".join(c for c in slug if c.isalnum() or c in "-_")
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "exemplars", f"{safe}.json")
    if not os.path.isfile(path):
        raise HTTPException(404, "Exemplar not found")
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


@router.get("/api/exemplars/{slug}/export.docx")
def export_exemplar(slug: str):
    import api_server
    return api_server.package_to_docx(get_exemplar(slug))


@router.get("/api/exemplars/{slug}/materials.docx")
def export_materials(slug: str):
    """The classroom-ready pack: worksheet, answer key and glossary only.

    Separate from the lesson plan on purpose. A teacher prints this and hands
    most of it out; they do not want the plan stapled to it.
    """
    import io

    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    import api_documents

    pkg = get_exemplar(slug)
    worksheet = pkg.get("worksheet") or {}
    doc = Document()
    doc.styles["Normal"].font.name = "Lato"
    doc.styles["Normal"].font.size = Pt(11)

    # --- Student worksheet -------------------------------------------------
    doc.add_heading(worksheet.get("title") or "Worksheet", level=0)
    doc.add_paragraph("Name: ______________________________    Class: ____________")
    if worksheet.get("learner_instructions"):
        doc.add_paragraph(worksheet["learner_instructions"])
    if worksheet.get("language_help_box"):
        doc.add_heading("Language help", level=2)
        for line in worksheet["language_help_box"]:
            doc.add_paragraph(str(line), style="List Bullet")
    for part in worksheet.get("parts", []):
        heading = re.sub(r"^\s*part\s*\d+\s*[.:\-\u2013]\s*", "", str(part.get("heading", "")),
                         flags=re.I)
        doc.add_heading(f"Part {part.get('number')}. {heading}", level=1)
        if part.get("instructions"):
            doc.add_paragraph(part["instructions"])
        for item in part.get("items", []):
            doc.add_paragraph(str(item.get("prompt", "")), style="List Number")
            for _ in range(max(1, int(item.get("lines") or 1))):
                doc.add_paragraph("_" * 78)
    if worksheet.get("exit_check"):
        doc.add_heading("Before you go", level=1)
        doc.add_paragraph(worksheet["exit_check"].get("prompt", ""))
        doc.add_paragraph("_" * 78)
        doc.add_paragraph("_" * 78)

    # --- Glossary, on its own page so it can be handed out separately ------
    if pkg.get("glossary"):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading("Glossary", level=0)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        headers = ["Term", "What it means", "Saying it",
                   f"In {pkg['header'].get('first_language') or 'your language'}"]
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
        for term in pkg["glossary"]:
            cells = table.add_row().cells
            cells[0].text = str(term.get("term", ""))
            cells[1].text = str(term.get("definition", ""))
            cells[2].text = str(term.get("pronunciation_note", ""))
            cells[3].text = ""

    # --- Answer key, teacher copy -----------------------------------------
    if pkg.get("answer_key"):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading("Answer key (teacher copy)", level=0)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells,
                              ["Part", "Item", "Answer", "Marking note"]):
            cell.text = text
        for answer in pkg["answer_key"]:
            cells = table.add_row().cells
            cells[0].text = str(answer.get("part", ""))
            cells[1].text = str(answer.get("prompt", ""))
            cells[2].text = str(answer.get("answer", ""))
            note = str(answer.get("marking_note", ""))
            also = str(answer.get("accept_also", ""))
            cells[3].text = f"{note}\nAlso accept: {also}" if also else note

    # --- Sources, because a handout without them is not defensible ---------
    sources = (pkg.get("grounded") or {}).get("sources", [])
    if sources:
        doc.add_heading("Sources for the subject content", level=1)
        for source in sources:
            doc.add_paragraph(f"[{source['index']}] {source['title']} — {source['url']}",
                              style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    name = api_documents.safe_filename(f"materials-{slug}", "docx")
    from fastapi.responses import StreamingResponse
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )
