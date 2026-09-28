"""
api_documents.py — routes for course-level updates.

Kept separate from api_server.py so the lesson pipeline and the reporting
pipeline can be read, tested and changed independently.
"""

from __future__ import annotations

import io
import json
import uuid
from datetime import datetime, timezone
from typing import Any

from fastapi import APIRouter, Body, Header, HTTPException
from fastapi.responses import StreamingResponse

import clil_reports

router = APIRouter()
_db = None


def attach(db) -> APIRouter:
    global _db
    _db = db
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS documents (
            id TEXT PRIMARY KEY,
            visitor_id TEXT NOT NULL,
            course_id TEXT NOT NULL,
            kind TEXT NOT NULL,
            share_token TEXT,
            data TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        );
        """
    )
    db.commit()
    return router


def vid(visitor: str | None) -> str:
    return visitor or "local-dev-visitor"


def safe_filename(stem: str, extension: str) -> str:
    """HTTP headers are latin-1 only, and titles come from teacher input, so
    reduce the name to a safe ASCII slug before it reaches a header."""
    slug = "".join(c if (c.isalnum() or c in "-_") else "-" for c in stem.encode("ascii", "ignore").decode())
    while "--" in slug:
        slug = slug.replace("--", "-")
    return f"{slug.strip('-')[:60] or 'document'}.{extension}"


def _course_bundle(course_id: str, visitor: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    row = _db.execute("SELECT id, data FROM courses WHERE id = ? AND visitor_id = ?",
                      [course_id, visitor]).fetchone()
    if not row:
        raise HTTPException(404, "Course not found")
    course = json.loads(row["data"])
    course["id"] = row["id"]
    rows = _db.execute(
        "SELECT id, package, reflection, created_at FROM lessons "
        "WHERE course_id = ? AND visitor_id = ? ORDER BY created_at",
        [course_id, visitor],
    ).fetchall()
    lessons = [
        {
            "id": r["id"],
            "package": json.loads(r["package"]),
            "reflection": json.loads(r["reflection"]) if r["reflection"] else None,
            "created_at": r["created_at"],
        }
        for r in rows
    ]
    return course, lessons


@router.post("/api/courses/{course_id}/documents", status_code=201)
def create_document(course_id: str, body: dict[str, Any] = Body(default={}),
                    x_visitor_id: str | None = Header(default=None)) -> dict[str, Any]:
    visitor = vid(x_visitor_id)
    kind = body.get("kind") or "newsletter"
    course, lessons = _course_bundle(course_id, visitor)
    if not lessons:
        raise HTTPException(400, "This course has no lessons yet, so there is nothing to report on.")

    prow = _db.execute("SELECT data FROM profiles WHERE visitor_id = ?", [visitor]).fetchone()
    profile = json.loads(prow["data"]) if prow else {}

    try:
        doc = clil_reports.build(kind, course, lessons, profile)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc

    doc["share_token"] = uuid.uuid4().hex
    doc["course_id"] = course_id
    _db.execute(
        "INSERT INTO documents (id, visitor_id, course_id, kind, share_token, data) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        [doc["id"], visitor, course_id, kind, doc["share_token"], json.dumps(doc)],
    )
    _db.commit()
    return doc


@router.get("/api/courses/{course_id}/documents")
def list_documents(course_id: str,
                   x_visitor_id: str | None = Header(default=None)) -> list[dict[str, Any]]:
    rows = _db.execute(
        "SELECT id, kind, data, created_at, updated_at FROM documents "
        "WHERE course_id = ? AND visitor_id = ? ORDER BY created_at DESC",
        [course_id, vid(x_visitor_id)],
    ).fetchall()
    out = []
    for r in rows:
        d = json.loads(r["data"])
        out.append({
            "id": r["id"],
            "kind": r["kind"],
            "kind_label": d.get("kind_label"),
            "title": d.get("title"),
            "period": d.get("period"),
            "created_at": r["created_at"],
            "updated_at": r["updated_at"],
        })
    return out


@router.get("/api/documents/{doc_id}")
def get_document(doc_id: str, token: str | None = None,
                 x_visitor_id: str | None = Header(default=None)) -> dict[str, Any]:
    row = _db.execute("SELECT * FROM documents WHERE id = ?", [doc_id]).fetchone()
    if not row:
        raise HTTPException(404, "Document not found")
    if row["visitor_id"] != vid(x_visitor_id) and token != row["share_token"]:
        raise HTTPException(403, "This document is not shared with you")
    return json.loads(row["data"])


@router.put("/api/documents/{doc_id}")
def update_document(doc_id: str, body: dict[str, Any] = Body(...),
                    x_visitor_id: str | None = Header(default=None)) -> dict[str, Any]:
    """Save the teacher's edits.

    Locked sections keep their generated wording. A teacher can delete a
    privacy or liability notice, but cannot silently rewrite it into something
    misleading, which is the behaviour that matters when a document leaves the
    building.
    """
    row = _db.execute("SELECT data FROM documents WHERE id = ? AND visitor_id = ?",
                      [doc_id, vid(x_visitor_id)]).fetchone()
    if not row:
        raise HTTPException(404, "Document not found")

    doc = json.loads(row["data"])
    edits = {s.get("id"): s for s in body.get("sections", [])}
    kept: list[dict[str, Any]] = []
    for section in doc["sections"]:
        edit = edits.get(section["id"])
        if edit is None:
            continue                                    # removed by the teacher
        if not section.get("locked"):
            section["heading"] = (edit.get("heading") or section["heading"]).strip()
            section["body"] = [str(x).strip() for x in (edit.get("body") or []) if str(x).strip()]
        kept.append(section)

    doc["sections"] = kept
    if body.get("title"):
        doc["title"] = str(body["title"]).strip()
    doc["edited_at"] = datetime.now(timezone.utc).isoformat(timespec="seconds")

    _db.execute("UPDATE documents SET data = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                [json.dumps(doc), doc_id])
    _db.commit()
    return doc


@router.delete("/api/documents/{doc_id}")
def delete_document(doc_id: str, x_visitor_id: str | None = Header(default=None)) -> dict[str, str]:
    _db.execute("DELETE FROM documents WHERE id = ? AND visitor_id = ?", [doc_id, vid(x_visitor_id)])
    _db.commit()
    return {"deleted": doc_id}


@router.get("/api/documents/{doc_id}/export.docx")
def export_document_docx(doc_id: str, token: str | None = None,
                         x_visitor_id: str | None = Header(default=None)):
    doc = get_document(doc_id, token, x_visitor_id)
    from docx import Document
    from docx.shared import Pt

    d = Document()
    d.styles["Normal"].font.name = "Lato"
    d.styles["Normal"].font.size = Pt(11)

    d.add_heading(doc["title"], level=0)
    d.add_paragraph(f"{doc['kind_label']} · for {doc['audience']} · {doc['period']}")
    d.add_paragraph(doc["from_line"])
    for key, value in doc["meta"]:
        if value:
            d.add_paragraph(f"{key}: {value}", style="List Bullet")

    for section in doc["sections"]:
        d.add_heading(section["heading"], level=1)
        for line in section["body"]:
            text = str(line)
            if text.startswith("• "):
                d.add_paragraph(text[2:], style="List Bullet")
            else:
                d.add_paragraph(text)

    d.add_heading(doc["notice"]["heading"], level=1)
    for line in doc["notice"]["body"]:
        d.add_paragraph(line)

    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)
    name = safe_filename(f"{doc['kind']}-{doc['title']}", "docx")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )
