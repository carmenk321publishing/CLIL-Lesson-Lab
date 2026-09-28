#!/usr/bin/env python3
"""
api_server.py — CLIL Lesson Lab backend. Runs on port 8000.

Responsibilities:
  - Teacher profile, courses, lessons and reflections (SQLite, visitor-scoped).
  - Package generation, delegated entirely to clil_engine.
  - Share links and exports (PDF via the browser, DOCX, PPTX).

Data is scoped by the X-Visitor-Id header the hosting proxy injects, because
sandboxed iframes have no cookies or localStorage. In production this would be
replaced by real teacher accounts.
"""

from __future__ import annotations

import io
import json
import os
import sqlite3
import uuid
from typing import Any

from fastapi import Body, FastAPI, Header, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse

import api_ai
import api_documents
import clil_engine
import clil_reports

DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data.db")
db = sqlite3.connect(DB_PATH, check_same_thread=False)
db.row_factory = sqlite3.Row

db.executescript(
    """
    CREATE TABLE IF NOT EXISTS profiles (
        visitor_id TEXT PRIMARY KEY,
        data TEXT NOT NULL,
        updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS courses (
        id TEXT PRIMARY KEY,
        visitor_id TEXT NOT NULL,
        data TEXT NOT NULL,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS lessons (
        id TEXT PRIMARY KEY,
        visitor_id TEXT NOT NULL,
        course_id TEXT,
        share_token TEXT,
        payload TEXT NOT NULL,
        package TEXT NOT NULL,
        reflection TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """
)
db.commit()

app = FastAPI(title="CLIL Lesson Lab API")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition"],
)


def vid(visitor: str | None) -> str:
    return visitor or "local-dev-visitor"


# ---------------------------------------------------------------------------
# Options and profile
# ---------------------------------------------------------------------------

app.include_router(api_documents.attach(db))
app.include_router(api_ai.attach(db))


@app.get("/api/options")
def options() -> dict[str, Any]:
    return {
        **clil_engine.OPTIONS,
        "cadences": clil_reports.CADENCES,
        "doc_kinds": clil_reports.DOC_KINDS,
    }


@app.get("/api/profile")
def get_profile(x_visitor_id: str | None = Header(default=None)) -> dict[str, Any]:
    row = db.execute("SELECT data FROM profiles WHERE visitor_id = ?", [vid(x_visitor_id)]).fetchone()
    return json.loads(row["data"]) if row else {}


@app.put("/api/profile")
def put_profile(body: dict[str, Any] = Body(...), x_visitor_id: str | None = Header(default=None)) -> dict[str, Any]:
    db.execute(
        "INSERT INTO profiles (visitor_id, data) VALUES (?, ?) "
        "ON CONFLICT(visitor_id) DO UPDATE SET data = excluded.data, updated_at = CURRENT_TIMESTAMP",
        [vid(x_visitor_id), json.dumps(body)],
    )
    db.commit()
    return body


# ---------------------------------------------------------------------------
# Courses
# ---------------------------------------------------------------------------

@app.get("/api/courses")
def list_courses(x_visitor_id: str | None = Header(default=None)) -> list[dict[str, Any]]:
    rows = db.execute(
        "SELECT id, data, created_at FROM courses WHERE visitor_id = ? ORDER BY created_at DESC",
        [vid(x_visitor_id)],
    ).fetchall()
    out = []
    for r in rows:
        course = json.loads(r["data"])
        course["id"] = r["id"]
        course["created_at"] = r["created_at"]
        lessons = db.execute(
            "SELECT id, package, reflection, created_at FROM lessons WHERE course_id = ? ORDER BY created_at",
            [r["id"]],
        ).fetchall()
        course["lessons"] = [
            {
                "id": l["id"],
                "title": json.loads(l["package"])["header"]["title"],
                "lesson_number": json.loads(l["package"])["header"].get("lesson_number"),
                "has_reflection": bool(l["reflection"]),
                "created_at": l["created_at"],
            }
            for l in lessons
        ]
        out.append(course)
    return out


@app.post("/api/courses", status_code=201)
def create_course(body: dict[str, Any] = Body(...), x_visitor_id: str | None = Header(default=None)) -> dict[str, Any]:
    cid = uuid.uuid4().hex[:12]
    db.execute("INSERT INTO courses (id, visitor_id, data) VALUES (?, ?, ?)",
               [cid, vid(x_visitor_id), json.dumps(body)])
    db.commit()
    body["id"] = cid
    body["lessons"] = []
    return body


@app.get("/api/courses/{course_id}")
def get_course(course_id: str, x_visitor_id: str | None = Header(default=None)) -> dict[str, Any]:
    row = db.execute("SELECT id, data FROM courses WHERE id = ? AND visitor_id = ?",
                     [course_id, vid(x_visitor_id)]).fetchone()
    if not row:
        raise HTTPException(404, "Course not found")
    course = json.loads(row["data"])
    course["id"] = row["id"]
    lessons = db.execute(
        "SELECT id, package, reflection, created_at FROM lessons WHERE course_id = ? ORDER BY created_at",
        [course_id],
    ).fetchall()
    course["lessons"] = [
        {
            "id": l["id"],
            "package": json.loads(l["package"]),
            "reflection": json.loads(l["reflection"]) if l["reflection"] else None,
            "created_at": l["created_at"],
        }
        for l in lessons
    ]
    return course


@app.delete("/api/courses/{course_id}")
def delete_course(course_id: str, x_visitor_id: str | None = Header(default=None)) -> dict[str, str]:
    db.execute("DELETE FROM courses WHERE id = ? AND visitor_id = ?", [course_id, vid(x_visitor_id)])
    db.execute("DELETE FROM lessons WHERE course_id = ? AND visitor_id = ?", [course_id, vid(x_visitor_id)])
    db.commit()
    return {"deleted": course_id}


# ---------------------------------------------------------------------------
# Generation
# ---------------------------------------------------------------------------

@app.post("/api/generate", status_code=201)
def generate(body: dict[str, Any] = Body(...), x_visitor_id: str | None = Header(default=None)) -> dict[str, Any]:
    course_id = body.get("course_id")
    if course_id:
        row = db.execute("SELECT data FROM courses WHERE id = ? AND visitor_id = ?",
                         [course_id, vid(x_visitor_id)]).fetchone()
        if row:
            course = json.loads(row["data"])
            # Course context is inherited, so the teacher never re-enters it.
            for key in ("subject", "target_language", "first_language", "learner_group",
                        "support_level", "purpose", "setting", "format", "teacher_role",
                        "minutes", "lesson_total"):
                body.setdefault(key, course.get(key))
            body.setdefault("course_name", course.get("name"))
            body["mode"] = "course"
            taught = db.execute("SELECT COUNT(*) c FROM lessons WHERE course_id = ?", [course_id]).fetchone()["c"]
            body.setdefault("lesson_number", taught + 1)

    package = clil_engine.generate(body)
    lid = package["id"]
    token = uuid.uuid4().hex
    db.execute(
        "INSERT INTO lessons (id, visitor_id, course_id, share_token, payload, package) VALUES (?, ?, ?, ?, ?, ?)",
        [lid, vid(x_visitor_id), course_id, token, json.dumps(body), json.dumps(package)],
    )
    db.commit()
    package["share_token"] = token
    return package


@app.get("/api/lessons/{lesson_id}")
def get_lesson(lesson_id: str, token: str | None = None,
               x_visitor_id: str | None = Header(default=None)) -> dict[str, Any]:
    row = db.execute("SELECT * FROM lessons WHERE id = ?", [lesson_id]).fetchone()
    if not row:
        raise HTTPException(404, "Lesson not found")
    if row["visitor_id"] != vid(x_visitor_id) and token != row["share_token"]:
        raise HTTPException(403, "This lesson package is not shared with you")
    pkg = json.loads(row["package"])
    pkg["share_token"] = row["share_token"]
    pkg["course_id"] = row["course_id"]
    return pkg


@app.get("/api/lessons")
def list_lessons(x_visitor_id: str | None = Header(default=None)) -> list[dict[str, Any]]:
    rows = db.execute(
        "SELECT id, course_id, package, reflection, created_at FROM lessons "
        "WHERE visitor_id = ? ORDER BY created_at DESC LIMIT 50",
        [vid(x_visitor_id)],
    ).fetchall()
    out = []
    for r in rows:
        p = json.loads(r["package"])
        out.append({
            "id": r["id"],
            "course_id": r["course_id"],
            "title": p["header"]["title"],
            "subject": p["header"]["subject"],
            "course_name": p["header"].get("course_name"),
            "lesson_number": p["header"].get("lesson_number"),
            "has_reflection": bool(r["reflection"]),
            "created_at": r["created_at"],
        })
    return out


# ---------------------------------------------------------------------------
# Reflection
# ---------------------------------------------------------------------------

@app.put("/api/lessons/{lesson_id}/reflection")
def save_reflection(lesson_id: str, body: dict[str, Any] = Body(...),
                    x_visitor_id: str | None = Header(default=None)) -> dict[str, Any]:
    row = db.execute("SELECT id FROM lessons WHERE id = ? AND visitor_id = ?",
                     [lesson_id, vid(x_visitor_id)]).fetchone()
    if not row:
        raise HTTPException(404, "Lesson not found")
    db.execute("UPDATE lessons SET reflection = ? WHERE id = ?", [json.dumps(body), lesson_id])
    db.commit()
    return body


@app.get("/api/lessons/{lesson_id}/reflection")
def get_reflection(lesson_id: str, x_visitor_id: str | None = Header(default=None)) -> dict[str, Any]:
    row = db.execute("SELECT reflection FROM lessons WHERE id = ? AND visitor_id = ?",
                     [lesson_id, vid(x_visitor_id)]).fetchone()
    if not row:
        raise HTTPException(404, "Lesson not found")
    return json.loads(row["reflection"]) if row["reflection"] else {}


@app.get("/api/courses/{course_id}/log")
def teaching_log(course_id: str, x_visitor_id: str | None = Header(default=None)) -> dict[str, Any]:
    """The accumulated reflections across a course, as a teaching log."""
    rows = db.execute(
        "SELECT id, package, reflection, created_at FROM lessons "
        "WHERE course_id = ? AND visitor_id = ? ORDER BY created_at",
        [course_id, vid(x_visitor_id)],
    ).fetchall()
    entries = []
    for r in rows:
        if not r["reflection"]:
            continue
        p = json.loads(r["package"])
        entries.append({
            "lesson_id": r["id"],
            "title": p["header"]["title"],
            "lesson_number": p["header"].get("lesson_number"),
            "date": r["created_at"],
            "reflection": json.loads(r["reflection"]),
        })
    return {"course_id": course_id, "entries": entries}


# ---------------------------------------------------------------------------
# Exports
# ---------------------------------------------------------------------------

def _flatten(pkg: dict[str, Any]) -> list[tuple[str, str, list[str]]]:
    """(heading_level, heading, paragraphs) tuples for document export."""
    h = pkg["header"]
    blocks: list[tuple[str, str, list[str]]] = []
    blocks.append(("h1", f"CLIL Lesson Package: {h['title']}", []))
    blocks.append(("meta", "", [
        f"Subject: {h['subject']}",
        f"Course: {h['course_name'] or 'Standalone lesson'}",
        f"Position: {pkg['sequence']['position']}",
        f"Learners: {h['learner_group']}  |  Purpose: {h['purpose']}",
        f"Target language: {h['target_language']}  |  First language: {h['first_language']}",
        f"Language support: {h['support_label']} ({h['support_cefr']})",
        f"Setting: {h['setting']}  |  Format: {h['format']}  |  Length: {h['minutes']} minutes",
        f"Planned by: {h['teacher_role']}",
    ]))
    blocks.append(("h2", pkg["sequence"]["heading"], pkg["sequence"]["carried_forward"]))

    obj = pkg["objectives"]
    blocks.append(("h2", f"Objectives (Bloom level: {obj['bloom_level']})", []))
    blocks.append(("h3", "Content objectives", [f"[{o['bloom']}] {o['text']}" for o in obj["content"]]))
    blocks.append(("h3", "Language objectives", [f"{o['strand']}: {o['text']}" for o in obj["language"]]))
    if obj.get("learning_skill"):
        blocks.append(("h3", "Learning skills objective", [obj["learning_skill"]]))

    for st in pkg["stages"]:
        blocks.append(("h2", f"Stage {st['number']}: {st['name']} ({st['minutes']} min)", [st["purpose"]]))
        blocks.append(("h3", "Teacher does", st["teacher_does"]))
        blocks.append(("h3", "Learners do", st["learners_do"]))
        blocks.append(("h3", "Language focus", [st["language_focus"]]))
        blocks.append(("h3", "Watch for", [st["watch_for"]]))
        if st.get("differentiation"):
            blocks.append(("h3", "Differentiation", st["differentiation"]))
        if st.get("evidence"):
            blocks.append(("h3", "Evidence of learning", st["evidence"]))

    lb = pkg["language_bank"]
    blocks.append(("h2", lb["heading"], [
        f"Text load: {lb['text_load']}",
        f"First language policy: {lb['first_language_policy']}",
        lb["vocabulary_note"],
    ]))
    blocks.append(("h3", "Functions", lb["functions"]))
    blocks.append(("h3", "Sentence frames", lb["frames"]))
    blocks.append(("h2", "Materials and preparation", pkg["materials"]))

    for blk in pkg["compensation"]["blocks"]:
        blocks.append(("h2", blk["heading"], [blk["why"]]))
        blocks.append(("h3", "Check these", blk["items"]))

    la = pkg["looking_ahead"]
    blocks.append(("h2", la["heading"], la["items"] + [la["reflection_prompt"]]))

    rc = pkg["review_checklist"]
    blocks.append(("h2", rc["heading"], []))
    for grp in rc["groups"]:
        blocks.append(("h3", grp["label"], grp["items"]))

    dc = pkg["disclaimer"]
    blocks.append(("h2", dc["heading"], dc["body"]))
    blocks.append(("h3", "Framework sources", [f"{s['label']} — {s['url']}" for s in dc["sources"]]))
    return blocks


@app.get("/api/lessons/{lesson_id}/export.docx")
def export_docx(lesson_id: str, token: str | None = None,
                x_visitor_id: str | None = Header(default=None)):
    return package_to_docx(get_lesson(lesson_id, token, x_visitor_id))


def package_to_docx(pkg: dict[str, Any]):
    """Render any lesson package as a Word document. Shared by lessons and exemplars."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Lato"
    style.font.size = Pt(11)

    for level, heading, paras in _flatten(pkg):
        if level == "h1":
            doc.add_heading(heading, level=0)
        elif level == "h2":
            doc.add_heading(heading, level=1)
        elif level == "h3":
            doc.add_heading(heading, level=2)
        for p in paras:
            if level in ("h3", "meta"):
                doc.add_paragraph(p, style="List Bullet")
            else:
                doc.add_paragraph(p)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    name = api_documents.safe_filename(f"CLIL-Lesson-Package-{pkg['header']['title']}", "docx")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


@app.get("/api/lessons/{lesson_id}/export.pptx")
def export_pptx(lesson_id: str, token: str | None = None,
                x_visitor_id: str | None = Header(default=None)):
    pkg = get_lesson(lesson_id, token, x_visitor_id)
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = pkg["header"]["title"] or "CLIL Lesson"
    title_slide.placeholders[1].text = (
        f"{pkg['header']['subject']}  |  {pkg['header']['learner_group']}\n"
        f"{pkg['sequence']['position']}  |  {pkg['header']['minutes']} minutes"
    )

    def bullet_slide(title: str, bullets: list[str]) -> None:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        tf = s.placeholders[1].text_frame
        tf.clear()
        tf.word_wrap = True
        for i, b in enumerate(bullets[:6]):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = b[:220]
            para.font.size = Pt(16)

    obj = pkg["objectives"]
    bullet_slide("Objectives", [f"[{o['bloom']}] {o['text']}" for o in obj["content"]]
                 + [f"{o['strand']}: {o['text']}" for o in obj["language"]])
    bullet_slide(pkg["sequence"]["heading"], pkg["sequence"]["carried_forward"])
    for st in pkg["stages"]:
        bullet_slide(f"{st['number']}. {st['name']} ({st['minutes']} min)",
                     [st["purpose"]] + st["teacher_does"][:2] + st["learners_do"][:2])
    bullet_slide(pkg["language_bank"]["heading"], pkg["language_bank"]["frames"])
    bullet_slide("Before you teach", [i for g in pkg["review_checklist"]["groups"] for i in g["items"]][:6])
    bullet_slide(pkg["disclaimer"]["heading"], pkg["disclaimer"]["body"])

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    name = api_documents.safe_filename(f"CLIL-Lesson-Package-{pkg['header']['title']}", "pptx")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


@app.get("/api/health")
def health() -> JSONResponse:
    return JSONResponse({"ok": True, "schema": clil_engine.SCHEMA_VERSION})


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000, log_level="info")
